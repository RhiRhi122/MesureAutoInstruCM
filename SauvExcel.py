import tkinter as tk
from tkinter import messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill
from datetime import datetime

# Fonction pour trouver la première colonne non fusionnée et vide
def find_first_unmerged_column(ws, row, start_col=3):
    col = start_col
    while True:
        cell = ws.cell(row=row, column=col)
        # Vérifie si la cellule fait partie d'une plage fusionnée
        is_merged = any(cell.coordinate in merged_cell for merged_cell in ws.merged_cells.ranges)
        if not is_merged and cell.value is None:  # Vérifie si la cellule n'est pas fusionnée et est vide
            return col
        col += 1

# Fonction pour sauvegarder les mesures dans un fichier Excel
def sauvegarder_mesures_excel(references_mesures, valeurs_mesures, etat_isolement, etat_mesure, operateur_nom, cryomodule_num, nom_poste, window):
    # Charger le fichier Excel
    fichier_excel = f"C:/Assemblage/LB650/Instrumentations/MesureAutoInstru/DevPython/TableauSondes{cryomodule_num}.xlsx"
    try:
        wb = openpyxl.load_workbook(fichier_excel)
    except FileNotFoundError:
        wb = openpyxl.Workbook()  # Créer un nouveau fichier si non trouvé
        ws = wb.active
        ws.append(["Référence", "Mesures"])  # Ajouter des en-têtes de colonne
    else:
        ws = wb.active

    # Trouver la première colonne vide en partant de la colonne C
    col = find_first_unmerged_column(ws, row=3)

    # Ajouter les informations générales (date, heure, opérateur, poste) dans la première colonne vide
    date_mesure = datetime.now().strftime("%d/%m/%Y")
    heure_mesure = datetime.now().strftime("%H:%M:%S")

    ws.cell(row=3, column=col, value=date_mesure)   # Date de la mesure
    ws.cell(row=4, column=col, value=heure_mesure)  # Heure de la mesure
    ws.cell(row=5, column=col, value=operateur_nom) # Nom de l'opérateur
    ws.cell(row=6, column=col, value=nom_poste)     # Nom du poste
    
    # Extraire la première valeur de chaque sous-liste
    valeurs_sauv = [sous_liste[0] for sous_liste in valeurs_mesures]

    # Boucle pour enregistrer les mesures avec les styles applicables
    for i, ref_mesure in enumerate(references_mesures):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            if row[0].value == ref_mesure:
                cell = ws.cell(row=row[0].row, column=col, value=valeurs_sauv[i])

                # Application des styles en fonction des états
                fill = None
                font = Font(bold=True)
                
                if not etat_isolement[i] and not etat_mesure[i]:
                    fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
                    font = Font(bold=True, color="FF0000")
                elif not etat_isolement[i]:
                    font = Font(bold=True, color="FF0000")
                elif not etat_mesure[i]:
                    fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
                    font = Font(bold=True, color="000000")

                if fill:
                    cell.fill = fill
                cell.font = font
                break
        else:
            # Si la référence n'est pas trouvée, ajoute une nouvelle ligne
            ws.append([ref_mesure] + [""] * (col - 3) + [valeurs_sauv[i]])

    # Fermer la fenêtre et sauvegarder le fichier
    window.destroy()
    wb.save(fichier_excel)
    messagebox.showinfo("Sauvegarde", "Les mesures ont été sauvegardées avec succès dans le fichier Excel.")
